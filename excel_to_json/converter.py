"""Core Excel -> JSON conversion. Importable from your own scripts:

    from excel_to_json.converter import convert_workbook
    data = convert_workbook("report.xlsx")  # {"Sheet1": [{...}, ...], ...}
"""
import datetime as dt
import json
from pathlib import Path

from openpyxl import load_workbook


def _clean(value):
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        return value.strip()
    return value


def _headers(row):
    """Unique, non-empty header names; blank ones become column_N."""
    seen, out = {}, []
    for i, cell in enumerate(row, 1):
        name = str(cell).strip() if cell not in (None, "") else f"column_{i}"
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 1
        out.append(name)
    return out


def sheet_names(path):
    wb = load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def convert_workbook(path, sheets=None, header_row=1, skip_empty_rows=True, drop_null_fields=False):
    """Return {sheet_name: [row_dict, ...]} for the selected sheets (default: all)."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        result = {}
        for name in sheets or wb.sheetnames:
            rows = wb[name].iter_rows(min_row=header_row, values_only=True)
            header = next(rows, None)
            if header is None:
                result[name] = []
                continue
            keys = _headers(header)
            records = []
            for row in rows:
                values = [_clean(v) for v in row]
                if skip_empty_rows and all(v in (None, "") for v in values):
                    continue
                rec = {k: (None if v == "" else v) for k, v in zip(keys, values)}
                if drop_null_fields:
                    rec = {k: v for k, v in rec.items() if v is not None}
                records.append(rec)
            result[name] = records
        return result
    finally:
        wb.close()


def to_json_text(data, indent=2):
    return json.dumps(data, indent=indent, ensure_ascii=False)


def write_json(data, out_path, indent=2):
    Path(out_path).write_text(to_json_text(data, indent), encoding="utf-8")
